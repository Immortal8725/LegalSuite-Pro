#!/usr/bin/env python3
"""Converts the firm-supplied source documents into Owlaw system templates.

PDF sources become a positioned JSON layout (text spans, vector drawings, images) so the
original layout is reproduced exactly; DOCX sources stay DOCX with their blanks turned into
{{field:original blank}} placeholders. Legal wording is never changed: an unfilled field
prints its original blank (underscores or [BRACKET] prompt).
"""
import base64, json, re, sys, pathlib
import fitz  # PyMuPDF
from docx import Document

SRC = pathlib.Path(sys.argv[1])
OUT = pathlib.Path(sys.argv[2])
OUT.mkdir(parents=True, exist_ok=True)

def snake(s):
    s = re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")
    return s[:40].strip("_")

# Context rules: (regex on text before the blank, key). First match wins.
RULES = [
    (r"(name\s*(&|and)\s*surname\)?\s*:?)\s*$", "client_name"),
    (r"identity\s*(number|document)?\s*(:|\))?\s*$", "client_id_number"),
    (r"id\s*no\s*:?\s*$", "client_id_number"),
    (r"(occurred|occurred on|that occurred on)\s*$", "accident_date"),
    (r"(signed|dated)\s*at\s*$", "signed_at"),
    (r"(on\s+this\s+the|on\s+this|this\s+the|this)\s*$", "signed_day"),
    (r"day\s+of\s*$", "signed_month"),
    (r"(^|\s|\}\})20$", "signed_year"),
    (r"attorney : \{\{attorney_name\}\}\s*$", "letter_day"),
    (r"\{\{letter_day[^}]*\}\}\s*$", "letter_month"),
    (r"terminate the mandate of\s*$", "previous_attorney"),
    (r"treated at\s*$", "hospital_name"),
    (r"saps\s*$", "saps_ar_number"),
    (r"instructions taken by\s*:?\s*$", "instructions_taken_by"),
    (r"^date\s*:\s*$", "instruction_date"),
    (r"(^|[^y])our ref\s*:\s*$", "file_number"),
    (r"your ref\s*:\s*$", "your_ref"),
    (r"estate late\s*:\s*$", "deceased_name"),
    (r"date of death\s*:\s*$", "date_of_death"),
    (r"surviving\s*$", "client_relationship"),
    (r"surviving\s+\{\{client_relationship[^}]*\}\}\s*$", "client_name"),
    (r"source of your accumulated wealth\?\s*$", "source_of_wealth"),
    (r"^1\.\s*$", "witness_1"),
    (r"^2\.\s*$", "witness_2"),
]
BLANK = re.compile(r"_{3,}")

class Namer:
    def __init__(self):
        self.generic = {}
    def name(self, before, section):
        b = before.strip().rstrip("_").strip()
        bl = b.lower()
        if section == "minor" and re.search(r"(name\s*(&|and)\s*surname\)?)\s*$", bl):
            return "represented_name"
        if section == "minor" and re.search(r"identity", bl[-30:]):
            return "represented_id_number"
        for rx, key in RULES:
            if re.search(rx, bl):
                return key
        words = re.sub(r"\{\{[^}]*\}\}", " ", b)
        words = re.findall(r"[A-Za-z]+", words)[-4:]
        key = snake(" ".join(words)) or "blank"
        n = self.generic.get(key, 0) + 1
        self.generic[key] = n
        return key if n == 1 else f"{key}_{n}"

def tokenise(text, namer, section, prefix="", whole=None):
    """Replace underscore blanks in text with {{key:____}} tokens, naming by context."""
    if re.fullmatch(r"[\s_]*_{3,}[\s_]*", whole if whole is not None else text):
        return text  # signature line: stays a physical line to sign on
    out, last = [], 0
    for m in BLANK.finditer(text):
        out.append(text[last:m.start()])
        before = prefix + "".join(out)
        if re.search(r"\}\}\s{10,}$", before):
            out.append(m.group(0))  # a second signature line on the same row
            last = m.end()
            continue
        key = namer.name(before, section)
        out.append("{{%s:%s}}" % (key, m.group(0)))
        last = m.end()
    out.append(text[last:])
    return "".join(out)

EMOJI_MAP = {"\U0001F4CD": "\u2302", "\U0001F4DE": "\u260E", "\u2709\ufe0f": "\u2709"}
EMOJI = re.compile("|".join(map(re.escape, EMOJI_MAP)))

def hexrgb(c):
    return "#%02x%02x%02x" % tuple(int(round(v * 255)) for v in c[:3])

def font_kind(name):
    n = name.lower()
    if "serif" in n or "times" in n or "tinos" in n or "gelasio" in n or "georgia" in n:
        return "serif"
    if "carlito" in n or "calibri" in n:
        return "carlito"
    return "sans"

def convert_pdf(path, pages, key, meta):
    doc = fitz.open(path)
    namer = Namer()
    out_pages = []
    section = ""
    for pno in pages:
        page = doc[pno - 1]
        d = page.get_text("dict")
        lines = []
        for block in d["blocks"]:
            if block["type"] == 1:
                continue
            for line in block["lines"]:
                spans = [s for s in line["spans"] if s["text"] != ""]
                if spans:
                    lines.append(spans)
        # merge lines that share a baseline so a row is processed left to right
        rows = {}
        for spans in lines:
            y = round(spans[0]["origin"][1], 0)
            k = next((r for r in rows if abs(r - y) < 2.5), y)
            rows.setdefault(k, []).extend(spans)
        out_lines = []
        prev_row = ""
        for y in sorted(rows):
            spans = sorted(rows[y], key=lambda s: s["origin"][0])
            linetext = ""
            whole = "".join(s["text"] for s in spans)
            prefix0 = (prev_row + " ") if whole.lstrip().startswith("_") else ""
            jspans = []
            for s in spans:
                t = s["text"]
                low = (linetext + t).lower()
                if "on behalf of" in low:
                    section = "minor"
                if "{{firm_logo}}" in t or re.search(r"^\s*(i, the undersigned|we, the undersigned)", low):
                    section = "" if "on behalf of" not in low else section
                tt = tokenise(t, namer, section, prefix0 + linetext, whole)
                linetext += tt
                f = s["flags"]
                jspans.append({
                    "x": round(s["origin"][0], 2), "y": round(s["origin"][1], 2),
                    "w": round(s["bbox"][2] - s["bbox"][0], 2),
                    "size": round(s["size"], 2), "font": font_kind(s["font"]),
                    "bold": bool(f & 16) or "bold" in s["font"].lower(),
                    "italic": bool(f & 2) or "italic" in s["font"].lower(),
                    "color": "#%06x" % s["color"], "text": EMOJI.sub(lambda m: EMOJI_MAP[m.group(0)], tt),
                    "bbox": [round(v, 2) for v in s["bbox"]], "invisible": s.get("alpha", 255) == 0})
            out_lines.append({"y": y, "spans": jspans})
            prev_row = linetext
        drawings = []
        for g in page.get_drawings():
            items = []
            for it in g["items"]:
                if it[0] == "l":
                    items.append(["l"] + [round(v, 2) for v in (it[1].x, it[1].y, it[2].x, it[2].y)])
                elif it[0] == "c":
                    items.append(["c"] + [round(v, 2) for pt in it[1:5] for v in (pt.x, pt.y)])
                elif it[0] == "re":
                    r = it[1]
                    items.append(["re", round(r.x0, 2), round(r.y0, 2), round(r.width, 2), round(r.height, 2)])
                elif it[0] == "qu":
                    q = it[1]
                    items.append(["qu"] + [round(v, 2) for pt in (q.ul, q.ur, q.lr, q.ll) for v in (pt.x, pt.y)])
            if not items:
                continue
            kind = g.get("type") or ""
            rect = g["rect"]
            drawings.append({
                "items": items, "bbox": [round(rect.x0, 2), round(rect.y0, 2), round(rect.x1, 2), round(rect.y1, 2)],
                "stroke": None if g.get("color") is None or "s" not in kind else hexrgb(g["color"]),
                "fill": None if g.get("fill") is None or "f" not in kind else hexrgb(g["fill"]),
                "lw": round(g.get("width") or 0.5, 2), "eo": bool(g.get("even_odd")),
                "close": bool(g.get("closePath"))})
        # Text drawn as outlines carries an invisible text layer. Keep the text (so fields can merge),
        # drop the outline glyphs, and draw the text visibly in the outline colour.
        for line in out_lines:
            for sp in line["spans"]:
                if not sp.pop("invisible", False):
                    continue
                bx = sp["bbox"]
                hits = [d for d in drawings if d["fill"] and d["bbox"][0] >= bx[0] - 3 and d["bbox"][2] <= bx[2] + 3
                        and d["bbox"][1] >= bx[1] - 6 and d["bbox"][3] <= bx[3] + 6]
                if hits:
                    sp["color"] = hits[0]["fill"]
                    if sp["size"] >= 14:
                        sp["bold"] = True
                    for d in hits:
                        drawings.remove(d)
                else:
                    sp["text"] = ""
        for line in out_lines:
            for sp in line["spans"]:
                sp.pop("bbox", None)
                sp.pop("invisible", None)
        images = []
        for info in page.get_image_info(xrefs=True):
            if not info.get("xref"):
                continue
            pix = fitz.Pixmap(doc, info["xref"])
            if pix.n - pix.alpha > 3:
                pix = fitz.Pixmap(fitz.csRGB, pix)
            b = info["bbox"]
            images.append({"x": round(b[0], 2), "y": round(b[1], 2), "w": round(b[2] - b[0], 2), "h": round(b[3] - b[1], 2),
                           "png": base64.b64encode(pix.tobytes("png")).decode()})
        out_pages.append({"w": round(page.rect.width, 2), "h": round(page.rect.height, 2),
                          "lines": out_lines, "drawings": drawings, "images": images})
    tpl = dict(meta, key=key, source="pdf", pages=out_pages)
    (OUT / f"{key}.json").write_text(json.dumps(tpl, ensure_ascii=False, separators=(",", ":")))
    return tpl

BRACKET = re.compile(r"\[([^\]\[]{2,80})\]")

def bracket_key(label, before, namer):
    l = label.lower()
    if "deceased" in l and "name" in l:
        return "deceased_name"
    if "estate number" in l or (l == "if available" and "estate number" in before.lower()):
        return "estate_number"
    if "date of death" in l:
        return "date_of_death"
    if "insert date" in l:
        return "letter_date"
    if "first child" in l:
        return "first_heir_name"
    if "second child" in l:
        return "second_heir_name"
    if l == "id number":
        return "second_heir_id_number" if "second" in before.lower() else "first_heir_id_number"
    if "attorney / law firm" in l:
        return "recipient_name"
    if l == "address":
        return "recipient_address"
    if l == "email address":
        return "recipient_email"
    return namer.name(label, "")

def convert_docx(path, key, meta):
    doc = Document(path)
    namer = Namer()
    section = ""
    for p in doc.paragraphs:
        text = p.text
        stripped = text.strip()
        if stripped and re.fullmatch(r"[0-9.]*\s*[A-Z][A-Z /&’'-]{3,}", stripped):
            section = snake(re.sub(r"^[0-9.]+", "", stripped)).replace("appointee_attorney_in_fact", "appointee")
        new = text
        if BRACKET.search(text):
            def rep(m):
                return "{{%s:[%s]}}" % (bracket_key(m.group(1), text[:m.start()], namer), m.group(1))
            new = BRACKET.sub(rep, text)
        new = tokenise(new, namer, section)
        m = re.fullmatch(r"(\s*)([A-Za-z][A-Za-z /&()-]{2,60}):\s*", new)
        if m and len(m.group(2).split()) > 4:
            m = None
        if m and section and "signature" not in m.group(2).lower():
            fk = snake(section + " " + m.group(2))
            new = new.rstrip() + " {{%s:____________________________}}" % fk
        elif m and "signature" in m.group(2).lower():
            new = new.rstrip() + " ____________________________"
        if new != text and p.runs:
            p.runs[0].text = new
            for r in p.runs[1:]:
                r.text = ""
    doc.save(OUT / f"{key}.docx")
    (OUT / f"{key}.meta.json").write_text(json.dumps(dict(meta, key=key, source="docx"), ensure_ascii=False))

A = SRC
f = lambda prefix: next(A.glob(prefix + "*"))
convert_pdf(f("44a1ac"), range(1, 16), "raf-special-power-of-attorney", {
    "name": "RAF claim pack: Special Power of Attorney and authorisations", "category": "Road Accident Fund",
    "description": "Special power of attorney, records authorisations, SAPS consent, payment instruction, termination of previous mandate, contingency fee agreement and instruction sheet."})
convert_pdf(f("0b8b2e"), range(1, 2), "new-practice-announcement", {
    "name": "New practice announcement", "category": "Correspondence",
    "description": "Letter to the profession and prospective clients announcing a new practice."})
convert_pdf(f("7b0f3a"), range(1, 3), "master-deceased-estate-report", {
    "name": "Letter to the Master: reporting a deceased estate", "category": "Deceased estates",
    "description": "Covering letter to the Master of the High Court lodging the estate documents."})
convert_pdf(f("6f1121"), range(1, 11), "client-mandate-fee-agreement", {
    "name": "Client mandate and fee agreement", "category": "Client onboarding",
    "description": "Mandate, fee agreement and Annexure B attorney and own client fee schedule."})
convert_pdf(f("6f1121"), range(11, 17), "privacy-policy", {
    "name": "Privacy policy (POPIA)", "category": "Client onboarding",
    "description": "Firm privacy policy with the information officer's contact details."})
convert_pdf(f("6f1121"), range(17, 19), "client-information-individual", {
    "name": "Client information sheet: individual", "category": "Client onboarding",
    "description": "FICA client information, verification and privacy consent for individuals, sureties, directors, members and trustees."})
convert_pdf(f("6f1121"), range(19, 21), "client-information-entity", {
    "name": "Client information sheet: company, CC or trust", "category": "Client onboarding",
    "description": "FICA client information, verification and privacy consent for companies, close corporations and trusts."})
convert_docx(f("8d799d"), "termination-of-mandate-deceased-estate", {
    "name": "Notice of termination of mandate (deceased estate)", "category": "Deceased estates",
    "description": "Heirs' notice terminating a previous attorney's mandate before Letters of Executorship are issued."})
convert_docx(f("450b1a"), "general-power-of-attorney-estate", {
    "name": "General power of attorney (estate of the late)", "category": "Deceased estates",
    "description": "Heirs appoint a representative to act on their behalf regarding the deceased estate."})
print("ok")
